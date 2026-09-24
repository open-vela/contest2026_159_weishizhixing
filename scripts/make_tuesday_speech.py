# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor


def font(run, size=12, bold=False, name="宋体", color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        font(run, 16 if level == 1 else 13, True, "黑体")


def para(doc, text, bold=False, size=12):
    p = doc.add_paragraph()
    run = p.add_run(text)
    font(run, size, bold)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.35


doc = Document()
for s in doc.sections:
    s.top_margin = Cm(2.2)
    s.bottom_margin = Cm(2.2)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.5)

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("见微·随身证")
font(r, 22, True, "黑体")

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("周二汇报口播稿（对照 PPT 逐页念）")
font(r, 16, True, "黑体")

t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("电脑旁请彭思羽同学看着网页。演示时把板子或鼠标交给同学，做完点头再往下讲。")
font(r, 11, False, "宋体", RGBColor(0x66, 0x66, 0x66))

para(doc, "用法：翻一页，先把「念」说完；到「演示」就停下来交给那位同学。括号里的字不用念。")

h(doc, "第 1 页　封面：见微 · 随身证")
para(doc, "念：各位老师好。我们组做的是「见微随身证」。它长得像挂在胸前的工牌。翻开是验真，合上是守护。我们不鉴定真假药，也不说官方认证，只帮老人先看清推销话术、先停手。下面我用两分钟把产品和现场演示过一遍。")

h(doc, "第 2 页　老人最怕的")
para(doc, "念：我们调研时发现，很多适老化产品是站在子女焦虑上做的，老人会觉得自己被盯着。真正高频的受骗场景，是眼前的包装、传单，和耳边的电话、讲座。诈骗最有效的手法，是把老人和真实信息隔开。老人往往还没起疑，就已经处在风险里。")

h(doc, "第 3 页　挂在胸前的校验点")
para(doc, "念：所以我们做的不是又一个聊天机器人，而是老人自己就能用的校验点。看得见的风险——包装和传单；听得到的风险——电话和讲座；说不出口的时候——连按三下，通知子女。不需要老人先学会复杂操作。")

h(doc, "第 4 页　怎么用：翻开 / 合上 / 三连按")
para(doc, "念：现场这块板子，用 BOOT 键代替翻盖。短按一下，表示翻开，进入验真；再短按一下，表示合上，进入守护；1.2 秒内连按三下，是求助。灯和网页上会用绿、橙、红告诉老人和子女现在是什么等级。接下来请同学们按这三步各演示一次。电脑页面请彭思羽同学帮忙看着，确认在线和结果。")

h(doc, "第 5 页　拍一拍普通包装")
para(doc, "念：先看验真。这是一瓶普通钙片，没有讲座话术。这里请杨立颖同学给大家演示一下拍照验真。立颖，把镜头对准瓶正面，短按一下 BOOT，等大约两秒让它自动拍。")
para(doc, "（等：屏上 VERIFY、AIM 2S、黄 PHOTO，然后绿 OK；网页出现「未见明显风险」，电脑可能会朗读。）")
para(doc, "念：好，板子是绿色的 OK，网页写「未见明显风险」。云端把图上的字读出来，再用同一套规则判断。我们没有说这药是真的，只是说眼前这句推销里，没有特别夸张的承诺。")

h(doc, "第 6 页　讲座里的推销话术")
para(doc, "念：对照一下。如果包装或传单上出现内部批文、天价、不让告诉子女这类话，应该打成高风险。为了现场稳定，我们用网页里的演示文案。这里请贾喻然同学给大家演示一下高风险验真。喻然，点「填入演示文案」，再点「开始验真」。这一步先不要按板子上的键。")
para(doc, "（等：网页大红「高风险」、先不要付款、电脑朗读。）")
para(doc, "念：刚才是绿的，现在是红的。变的是话术，不是我们宣布这瓶药是假药。建议是先不要付款，问清再买，并告诉子女。")

h(doc, "第 7 页　听一听电话话术")
para(doc, "念：合上证件，进入守护。板子听几秒电话内容，云端听写，还是同一套规则。这里请马凤玲同学给大家演示一下听音守护。凤玲，再短按一下 BOOT 切到守护，屏上出现 LISTEN 时，把手机免提贴紧麦孔，大声说：公安局让你转到安全账户，不要告诉家人。")
para(doc, "（等约 6 秒：板子红 HIGH；网页高风险、96110。若听不清出现橙 CHECK，改口：没贴紧时会提示没听清，请喻然用网页「填入演示话术」再走一遍。）")
para(doc, "念：命中公安局、安全账户、不让告诉家人，就是高风险。步骤是安抚、挂断后用通讯录回拨、不要转到所谓安全账户，必要时打 96110。")

h(doc, "第 8 页　连按三下")
para(doc, "念：老人当时可能说不出口。这里请袁艺轩同学给大家演示一下求助。艺轩，请在大约一秒内连按 BOOT 三下。按不准的话，请彭思羽同学点网页上的「模拟三连按求助」。")
para(doc, "（等：板子 HELP；网页求助变为「是」、红点。）")
para(doc, "念：子女侧立刻能看到求助。演示结束可以点一下「清除求助」。")

h(doc, "第 9 页　同一套规则")
para(doc, "念：整条链路可以分成三块。板子负责拍、听、按，以及把等级用大色块显示出来。对错不在板子里写死，而在云端同一套 Skill。子女看板同步结果，并用电脑朗读，因为这块开发板没有喇叭。贾喻然同学负责云端和知识库，杨立颖、马凤玲分别对接拍照和听音，袁艺轩做模式切换和状态机，彭思羽负责把板子和电脑联在同一网络上。")

h(doc, "第 10 页　Arduino 已通 / openvela 下一步")
para(doc, "念：今天现场用的是 Arduino 回退固件，目的是先把产品三条闭环跑稳。赛题要求基于 openvela。板端应用已经按同一份 HTTP 协议写在仓库里，还没有在 Ubuntu 上编烧到这块板。所以向老师汇报的现状是：基础功能可以演示；和 openvela 的真机对接，是下一步，不是还没做产品。")

h(doc, "第 11 页　谢谢老师")
para(doc, "念：我们希望老人在没有子女陪同时，胸前仍有一个温和的校验点：先停手，再问清。以上是见微随身证的汇报，谢谢老师，欢迎提问。")

h(doc, "提问时怎么接（不用投影，备用）")
para(doc, "若问真假药：我们只看话术和公开规则，不做药检，也不说国家认证。")
para(doc, "若问为什么不是 openvela：现场为了稳定先用 Arduino；openvela 应用与协议已经对齐，尚未烧录。")
para(doc, "若点到个人：只说自己做了什么、还没做什么，不要替别人展开。")

out = Path(r"C:\Users\jwr66\Documents\jianwei-id\docs\见微随身证_周二汇报口播稿.docx")
doc.save(out)
print(out)
