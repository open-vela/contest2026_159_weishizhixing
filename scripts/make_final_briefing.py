# -*- coding: utf-8 -*-
"""最终汇报：PPT + 与页码一一对应的口播稿。"""
from __future__ import annotations

import shutil
from pathlib import Path

from PIL import Image, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn as wqn
from docx.shared import Cm, Pt, RGBColor as DocRGB
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt as PptPt

ROOT = Path(r"C:\Users\jwr66\Documents\jianwei-id")
ASSETS = Path(r"C:\Users\jwr66\.cursor\projects\c-Users-jwr66-Documents-jianwei-id") / "assets"
PHOTO = ROOT / "docs" / "ppt_photos"
PHOTO.mkdir(parents=True, exist_ok=True)

W, H = Inches(13.333), Inches(7.5)
TOTAL = 13

CREAM = RGBColor(0xFB, 0xF4, 0xEC)
PAPER = RGBColor(0xFF, 0xFB, 0xF6)
INK = RGBColor(0x4A, 0x3A, 0x30)
SOFT = RGBColor(0x8A, 0x73, 0x64)
BROWN = RGBColor(0xA0, 0x6B, 0x52)
CORAL = RGBColor(0xC4, 0x6B, 0x5A)
SAGE = RGBColor(0x6A, 0x8F, 0x72)
GOLD = RGBColor(0xC9, 0xA0, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xE8, 0xDC, 0xCE)
HEAD = RGBColor(0xF3, 0xE6, 0xD8)

SRC = {
    "yao": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_image-a39beb5a-9f46-4105-857f-3c4be378ec97.png",
    "taisu": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_image-af1c6ce5-b21b-4676-a5b0-d5ffe9c6e212.png",
    "check": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_192d0201aa0f35ffff42cc52717a818e-7bf8af0f-88c6-4920-801c-841815b04321.jpg",
    "aim": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_aaae07cfaab53fb7327377c0149c2b64-2f0b1afa-5923-4b74-bfd7-99fcec642fce.jpg",
    "talk": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_def432a17fc9b8e270cc5f369ec14890-b554e09b-6813-4b7e-b209-d7cac8f11e53.jpg",
    "wifi": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_image-b9f80a45-d861-4446-8058-3081e83104b5.png",
    "beat": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_image-190c8b5c-e77f-4133-a8d0-290f92c56596.png",
    "guard_web": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_image-c8c34865-3efa-4b35-bb19-6ba18600d55e.png",
    "high": "c__Users_jwr66_AppData_Roaming_Cursor_User_workspaceStorage_4cfabee12cee61bdc0f5bca0b28a1909_images_8b1fdef7cc263bbf68ec67d0031b5ae4-b6891332-b66c-48f4-bb80-90df45139642.jpg",
}


def collect_photos() -> dict[str, Path]:
    out = {}
    for key, name in SRC.items():
        src = ASSETS / name
        if not src.exists():
            raise FileNotFoundError(src)
        dst = PHOTO / f"{key}{src.suffix.lower()}"
        shutil.copy2(src, dst)
        im = ImageOps.exif_transpose(Image.open(dst))
        if im.mode not in ("RGB", "RGBA"):
            im = im.convert("RGB")
        im.save(dst)
        out[key] = dst
    return out


def set_run(run, size, bold=False, color=INK, font="微软雅黑"):
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", font)


def fill_slide(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, l, t, w, h, color, radius=0.08):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.adjustments[0] = radius
    return shape


def add_text(slide, l, t, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color)
    return box


def add_lines(slide, l, t, w, h, lines, size, color=INK, gap=6, bold=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = PptPt(gap)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return box


def footer(slide, page):
    add_text(slide, Inches(0.55), Inches(7.12), Inches(7.5), Inches(0.28), "见微 · 随身证", 11, color=SOFT)
    add_text(
        slide,
        Inches(10.4),
        Inches(7.12),
        Inches(2.4),
        Inches(0.28),
        f"第{page}页",
        12,
        True,
        BROWN,
        PP_ALIGN.RIGHT,
    )


def accent_bar(slide):
    add_rect(slide, Inches(0), Inches(0), W, Inches(0.1), BROWN, 0)


def fit_pic(slide, path, l, t, w, h):
    im = Image.open(path)
    ar = im.width / float(im.height)
    box_w = float(w)
    box_h = float(h)
    box_ar = box_w / box_h
    if ar > box_ar:
        nh = int(box_w / ar)
        top = int(t) + (int(box_h) - nh) // 2
        slide.shapes.add_picture(str(path), int(l), top, int(box_w), nh)
    else:
        nw = int(box_h * ar)
        left = int(l) + (int(box_w) - nw) // 2
        slide.shapes.add_picture(str(path), left, int(t), nw, int(box_h))


def framed_pic(slide, path, l, t, w, h):
    add_rect(slide, l, t, w, h, PAPER, 0.06)
    pad = Inches(0.08)
    fit_pic(slide, path, l + pad, t + pad, w - pad * 2, h - pad * 2)


def set_cell(cell, text, size=12, bold=False, color=INK, fill=None, align=PP_ALIGN.LEFT):
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill
    else:
        cell.fill.solid()
        cell.fill.fore_color.rgb = PAPER
    cell.text = text
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    for p in tf.paragraphs:
        p.alignment = align
        for run in p.runs:
            set_run(run, size, bold, color)


def make_ppt(photos: dict[str, Path]) -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 封面
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    add_rect(s, Inches(3.55), Inches(1.05), Inches(6.25), Inches(3.35), PAPER, 0.1)
    add_rect(s, Inches(3.55), Inches(1.05), Inches(6.25), Inches(0.16), BROWN, 0)
    add_text(s, Inches(3.6), Inches(1.4), Inches(6.15), Inches(0.35), "挂在胸前的校验点", 15, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(3.6), Inches(1.85), Inches(6.15), Inches(0.85), "见微", 52, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(3.6), Inches(2.7), Inches(6.15), Inches(0.5), "随  身  证", 26, True, BROWN, PP_ALIGN.CENTER)
    add_text(s, Inches(3.6), Inches(3.35), Inches(6.15), Inches(0.4), "翻开验真  ·  合上守护  ·  三连按求助", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2), Inches(4.7), Inches(9.3), Inches(0.45), "先看清话术，先停手", 22, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(2), Inches(5.3), Inches(9.3), Inches(0.35), "不做药检，不谈官方认证，只帮老人在付款前多停一步", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(2), Inches(5.85), Inches(9.3), Inches(0.35), "汇报人  袁艺轩", 16, True, BROWN, PP_ALIGN.CENTER)
    add_text(s, Inches(2), Inches(6.25), Inches(9.3), Inches(0.3), "杨立颖  ·  马凤玲  ·  袁艺轩  ·  贾喻然  ·  彭思羽", 13, color=SOFT, align=PP_ALIGN.CENTER)
    footer(s, 1)

    # 2 场景
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 2)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.35), "我们为什么做它", 14, color=SOFT)
    add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.7), "老人最怕的，不是不会用手机", 28, True, INK)
    add_text(s, Inches(0.7), Inches(1.55), Inches(12), Inches(0.55), "是被隔开真实信息之后，还来不及起疑", 22, True, BROWN)
    cards = [
        (0.7, "眼前", "包装、传单、讲座材料\n字很大，承诺也很大"),
        (4.85, "耳边", "电话、免提、催转账\n越急，越不让告诉家人"),
        (9.0, "说不出口", "旁边坐着推销的人\n老人需要一个安静的出口"),
    ]
    for x, title, body in cards:
        add_rect(s, Inches(x), Inches(2.4), Inches(3.6), Inches(3.9), PAPER, 0.1)
        add_text(s, Inches(x), Inches(2.7), Inches(3.6), Inches(0.55), title, 22, True, BROWN, PP_ALIGN.CENTER)
        add_lines(s, Inches(x + 0.25), Inches(3.45), Inches(3.1), Inches(2.5), body.split("\n"), 15, SOFT, gap=10)

    # 3 产品
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 3)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.35), "它是什么", 14, color=SOFT)
    add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.65), "胸前的证件卡，不是又一个聊天窗口", 26, True, INK)
    add_lines(
        s,
        Inches(0.7),
        Inches(1.65),
        Inches(12),
        Inches(1.3),
        [
            "老人不用先学会复杂操作，翻开、合上、连按，三步就够。",
            "板子负责拍、听、按，并把等级用大色块告诉人；对错不在板子里写死。",
            "子女侧有一份同步看板，方便事后看清刚才发生了什么。",
        ],
        16,
        INK,
        gap=8,
    )
    trio = [
        (0.7, "看得见", "拍包装 / 传单", "验真"),
        (4.85, "听得到", "听电话 / 讲座", "守护"),
        (9.0, "说不出", "1.2 秒内连按三下", "求助"),
    ]
    for x, a, b, c in trio:
        add_rect(s, Inches(x), Inches(3.35), Inches(3.6), Inches(3.0), PAPER, 0.1)
        add_text(s, Inches(x), Inches(3.55), Inches(3.6), Inches(0.45), a, 14, color=SOFT, align=PP_ALIGN.CENTER)
        add_text(s, Inches(x), Inches(4.05), Inches(3.6), Inches(0.7), b, 20, True, INK, PP_ALIGN.CENTER)
        add_text(s, Inches(x), Inches(4.85), Inches(3.6), Inches(0.7), c, 24, True, BROWN, PP_ALIGN.CENTER)

    # 4 用法
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 4)
    add_text(s, Inches(0.7), Inches(0.4), Inches(12), Inches(0.35), "现场这块板怎么用", 14, color=SOFT)
    add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.6), "BOOT 键暂时代替翻盖", 26, True, INK)
    add_text(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.4), "绿是未见明显风险，橙是需要再核实，红是先停手。我们看的是话术，不是宣布药品真假。", 15, color=SOFT)
    rows = [
        (2.15, BROWN, "短按一下", "翻开 · 验真", "对准包装，约两秒后自动拍"),
        (3.7, SAGE, "再短按一下", "合上 · 守护", "对着麦说完，再按一次结束"),
        (5.25, CORAL, "连按三下", "求助 · 通知家人", "说不出口时，子女侧立刻看到"),
    ]
    for y, c, a, b, d in rows:
        add_rect(s, Inches(0.7), Inches(y), Inches(11.9), Inches(1.35), PAPER, 0.08)
        add_rect(s, Inches(0.7), Inches(y), Inches(0.16), Inches(1.35), c, 0)
        add_text(s, Inches(1.15), Inches(y + 0.35), Inches(3.2), Inches(0.65), a, 22, True, INK)
        add_text(s, Inches(4.5), Inches(y + 0.22), Inches(3.6), Inches(0.5), b, 20, True, c)
        add_text(s, Inches(4.5), Inches(y + 0.75), Inches(7.6), Inches(0.4), d, 14, color=SOFT)

    # 5 板子状态
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 5)
    add_text(s, Inches(0.7), Inches(0.35), Inches(12), Inches(0.3), "老人低头就能看懂", 14, color=SOFT)
    add_text(s, Inches(0.7), Inches(0.7), Inches(12), Inches(0.5), "屏上只有大字和颜色，不堆说明书", 24, True, INK)
    add_text(s, Inches(0.7), Inches(1.25), Inches(12), Inches(0.35), "翻开后先对准；合上后先开口。灯亮表示进入验真。", 14, color=SOFT)
    framed_pic(s, photos["aim"], Inches(0.6), Inches(1.7), Inches(6.0), Inches(4.7))
    framed_pic(s, photos["talk"], Inches(6.75), Inches(1.7), Inches(6.0), Inches(4.7))
    add_text(s, Inches(0.6), Inches(6.5), Inches(6.0), Inches(0.32), "VERIFY  ·  AIM 2S", 13, True, BROWN, PP_ALIGN.CENTER)
    add_text(s, Inches(6.75), Inches(6.5), Inches(6.0), Inches(0.32), "GUARD  ·  TALK", 13, True, SAGE, PP_ALIGN.CENTER)

    # 6 验真读字
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 6)
    add_text(s, Inches(0.55), Inches(0.32), Inches(6.2), Inches(0.28), "验真 · 读包装上的字", 14, color=SOFT)
    add_text(s, Inches(0.55), Inches(0.62), Inches(6.2), Inches(0.55), "先抄下来，再套同一套规则", 22, True, INK)
    add_lines(
        s,
        Inches(0.55),
        Inches(1.25),
        Inches(5.9),
        Inches(1.6),
        [
            "板子拍照传到云端，把图上的字填进看板。",
            "普通药膏盒上主要是品名，没有「包治、内部批文、不让告诉子女」。",
            "结论会写成「未见明显风险」——意思是眼前这句推销不夸张，不是给这支药做认证。",
        ],
        14,
        INK,
        gap=6,
    )
    framed_pic(s, photos["yao"], Inches(6.55), Inches(0.45), Inches(6.25), Inches(4.55))
    framed_pic(s, photos["beat"], Inches(0.55), Inches(3.05), Inches(5.85), Inches(3.55))

    # 7 验真话术
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 7)
    add_text(s, Inches(0.55), Inches(0.32), Inches(6.3), Inches(0.28), "验真 · 对照传单话术", 14, color=SOFT)
    add_text(s, Inches(0.55), Inches(0.62), Inches(6.3), Inches(0.5), "专家、包治、全家遭殃，就会升高风险", 20, True, INK)
    add_lines(
        s,
        Inches(0.55),
        Inches(1.2),
        Inches(6.15),
        Inches(5.3),
        [
            "同一套规则，换一张讲座传单。",
            "图上能读到：教授头衔、国家中医专家、治病除根、三个月行走自如、中风十年。",
            "这类材料不是让老人「再研究研究」，而是先不要付款、先问子女。",
            "绿和红对照的是话术密度，不是我们给产品盖章。",
            "子女看板上会留下读到的原文，方便事后核对。",
        ],
        15,
        INK,
        gap=10,
    )
    framed_pic(s, photos["taisu"], Inches(6.85), Inches(0.4), Inches(5.95), Inches(6.5))

    # 8 守护高风险
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 8)
    add_text(s, Inches(0.55), Inches(0.32), Inches(12), Inches(0.28), "守护 · 电话话术", 14, color=SOFT)
    add_text(s, Inches(0.55), Inches(0.62), Inches(12), Inches(0.5), "公安局、安全账户、不要告诉儿子 —— 先停手", 20, True, INK)
    add_text(s, Inches(0.55), Inches(1.15), Inches(12), Inches(0.35), "合上证件开始听。听清后走和验真同一套关键词与组合规则。板子整屏变红，网页给出安抚和回拨建议。", 14, color=SOFT)
    framed_pic(s, photos["guard_web"], Inches(0.5), Inches(1.6), Inches(7.55), Inches(5.25))
    framed_pic(s, photos["high"], Inches(8.2), Inches(1.6), Inches(4.6), Inches(5.25))

    # 9 CHECK
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 9)
    add_text(s, Inches(0.55), Inches(0.35), Inches(6.4), Inches(0.28), "听不清的时候", 14, color=SOFT)
    add_text(s, Inches(0.55), Inches(0.7), Inches(6.4), Inches(0.55), "宁可变橙，也不假装听懂", 22, True, INK)
    add_lines(
        s,
        Inches(0.55),
        Inches(1.4),
        Inches(6.3),
        Inches(5.0),
        [
            "麦在板子侧面的小孔上，免提要贴紧、大声说完再按 BOOT。",
            "没贴紧、环境太吵、录音太短，屏上是橙底 CHECK，网页写「没听清」。",
            "这是有意为之：不清楚就不给绿灯，避免老人被一句「没事」放过去。",
            "可以再听一次，也可以让子女在看板上改完文字再判断。",
        ],
        15,
        INK,
        gap=10,
    )
    framed_pic(s, photos["check"], Inches(7.05), Inches(0.55), Inches(5.7), Inches(6.2))

    # 10 架构
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 10)
    add_text(s, Inches(0.7), Inches(0.38), Inches(12), Inches(0.28), "对错写在同一处", 14, color=SOFT)
    add_text(s, Inches(0.7), Inches(0.72), Inches(12), Inches(0.5), "板子感知，云端 Skill 判断，子女看板同步", 22, True, INK)
    trio = [
        (0.7, "板子", "拍、听、按\n大色块 HIGH / CHECK / OK"),
        (4.85, "云端", "读图、听写、规则库\nskills/ 里的话术与步骤"),
        (9.0, "看板", "原文、等级、建议\n方便子女事后看"),
    ]
    for x, a, b in trio:
        add_rect(s, Inches(x), Inches(1.45), Inches(3.6), Inches(2.55), PAPER, 0.1)
        add_text(s, Inches(x), Inches(1.6), Inches(3.6), Inches(0.5), a, 20, True, BROWN, PP_ALIGN.CENTER)
        add_lines(s, Inches(x + 0.25), Inches(2.2), Inches(3.1), Inches(1.5), b.split("\n"), 14, INK, gap=6)
    add_text(s, Inches(0.7), Inches(4.15), Inches(12), Inches(0.35), "板子连上以后，心跳会持续把模式和最近一次结论带回来。", 14, color=SOFT)
    framed_pic(s, photos["beat"], Inches(2.4), Inches(4.55), Inches(8.5), Inches(2.35))

    # 11 表格
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 11)
    add_text(s, Inches(0.55), Inches(0.32), Inches(12), Inches(0.28), "分工对照", 14, color=SOFT)
    add_text(s, Inches(0.55), Inches(0.62), Inches(12), Inches(0.45), "已经做成的，和还可以精进的", 22, True, INK)
    add_text(s, Inches(0.55), Inches(1.08), Inches(12), Inches(0.32), "给老师看进度。细处不必展开，现场一句带过即可。", 13, color=SOFT)

    table = s.shapes.add_table(6, 3, Inches(0.45), Inches(1.5), Inches(12.4), Inches(5.35)).table
    table.columns[0].width = Inches(1.7)
    table.columns[1].width = Inches(5.5)
    table.columns[2].width = Inches(5.2)
    headers = ["同学", "已经做成", "还可以精进"]
    for i, h in enumerate(headers):
        set_cell(table.cell(0, i), h, 13, True, WHITE, BROWN, PP_ALIGN.CENTER)
    rows = [
        ("杨立颖", "板载拍照、包装图上传、读图进看板", "选更清晰的帧、补光、openvela 摄像头真机"),
        ("马凤玲", "按键结束录音、听写、听不清给橙灯", "贴麦体验、方言、省电唤醒"),
        ("袁艺轩", "验真 / 守护 / 求助状态机、色块等级、心跳", "真翻盖开关、板上独立路由"),
        ("贾喻然", "云端规则与 Skill、子女看板、本机读图听写", "评估集回归、板载真正发声"),
        ("彭思羽", "同一网络联调、协议对齐、端到端跑通", "热点切换、外壳、参赛五分钟视频"),
    ]
    for r, (name, done, todo) in enumerate(rows, 1):
        bg = PAPER if r % 2 else HEAD
        set_cell(table.cell(r, 0), name, 13, True, BROWN, bg, PP_ALIGN.CENTER)
        set_cell(table.cell(r, 1), done, 12, False, INK, bg)
        set_cell(table.cell(r, 2), todo, 12, False, INK, bg)

    # 12 现在与下一步
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 12)
    add_text(s, Inches(0.7), Inches(0.45), Inches(12), Inches(0.3), "向老师交代的现状", 14, color=SOFT)
    add_text(s, Inches(0.7), Inches(0.9), Inches(12), Inches(0.55), "产品三条闭环已经跑通，赛题真机还在下一步", 24, True, INK)
    add_rect(s, Inches(0.7), Inches(1.8), Inches(5.85), Inches(4.55), PAPER, 0.1)
    add_text(s, Inches(0.7), Inches(2.05), Inches(5.85), Inches(0.4), "已经做成", 14, True, SAGE, PP_ALIGN.CENTER)
    add_lines(
        s,
        Inches(1.05),
        Inches(2.6),
        Inches(5.2),
        Inches(3.4),
        [
            "ESP32-S3-EYE + Arduino 固件",
            "验真拍照、守护听音、三连按求助",
            "云端同一套 Skill，网页同步等级",
            "绿 / 橙 / 红，老人低头能看懂",
        ],
        16,
        INK,
        gap=12,
    )
    add_rect(s, Inches(6.8), Inches(1.8), Inches(5.85), Inches(4.55), PAPER, 0.1)
    add_text(s, Inches(6.8), Inches(2.05), Inches(5.85), Inches(0.4), "接下来", 14, True, GOLD, PP_ALIGN.CENTER)
    add_lines(
        s,
        Inches(7.15),
        Inches(2.6),
        Inches(5.2),
        Inches(3.4),
        [
            "openvela 真机编烧（协议已对齐）",
            "外壳与真翻盖，而不是 BOOT 代替",
            "听写再贴紧、再稳一些",
            "参赛仓与五分钟视频，截止日期 9 月 20 日",
        ],
        16,
        INK,
        gap=12,
    )

    # 13 谢谢
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 13)
    add_text(s, Inches(1), Inches(2.15), Inches(11.3), Inches(0.9), "先停手，再问清", 36, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(3.2), Inches(11.3), Inches(0.5), "谢谢老师", 22, True, BROWN, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(4.0), Inches(11.3), Inches(0.4), "欢迎提问", 16, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(5.35), Inches(11.3), Inches(0.35), "汇报人  袁艺轩", 14, color=BROWN, align=PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(5.8), Inches(11.3), Inches(0.35), "杨立颖  ·  马凤玲  ·  袁艺轩  ·  贾喻然  ·  彭思羽", 14, color=SOFT, align=PP_ALIGN.CENTER)

    out = ROOT / "docs" / "见微随身证_最终汇报.pptx"
    prs.save(out)
    return out


def font(run, size=12, bold=False, name="宋体", color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(wqn("w:eastAsia"), name)
    if color:
        run.font.color.rgb = color


def h(doc, text):
    p = doc.add_heading(text, level=1)
    for run in p.runs:
        font(run, 16, True, "黑体")


def para(doc, text, size=12, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    font(run, size, bold)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.38


def make_speech() -> Path:
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("见微·随身证")
    font(r, 22, True, "黑体")

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("最终汇报口播稿（对照 PPT 第几页念第几页）")
    font(r, 15, True, "黑体")

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("汇报人：袁艺轩　　不请同学上台演示，截图已经在片子里")
    font(r, 11, False, "宋体", DocRGB(0x66, 0x66, 0x66))

    para(doc, "用法：翻到第 N 页，再念「第 N 页」。括号里的字给自己看，不要念。全场约四到五分钟。不要说鉴定真假药，不要说已经在 openvela 真机上跑通。")

    h(doc, "第1页　封面")
    para(doc, "念：各位老师好。我是袁艺轩。我们组做的是「见微随身证」，形状就是挂在胸前的证件卡。翻开是验真，合上是守护，说不出口时连按三下求助。我们不做药检，也不谈官方认证，只帮老人先看清推销话术、先停手。今天用已经跑通的记录给老师看，不再单独做现场演示。")

    h(doc, "第2页　老人最怕什么")
    para(doc, "念：很多适老化产品，其实是站在子女焦虑上做的，老人会觉得自己被盯着。我们调研时更常碰到的，是三类现场：眼前的包装和传单，耳边的电话和讲座，以及旁边坐着人、老人不好意思开口。诈骗最有效的手法，是把老人和真实信息隔开。往往还没起疑，人已经处在风险里。")

    h(doc, "第3页　它是什么")
    para(doc, "念：所以我们做的不是又一个聊天机器人，而是老人自己胸前就能用的校验点。看得见的风险就拍；听得到的风险就听；说不出口就连按三下通知子女。板子负责感知和把等级用大色块显示出来，对错不在板子里写死，子女侧还有一份看板，方便事后看清刚才发生了什么。")

    h(doc, "第4页　BOOT 怎么用")
    para(doc, "念：正式产品想做成翻盖。现在这块开发板还没有翻盖微动开关，所以用 BOOT 键代替。短按一下，表示翻开，进入验真；再短按，表示合上，进入守护；大约一秒内连按三下，是求助。绿是未见明显风险，橙是需要再核实，红是先停手。请老师记住：变色看的是话术，不是我们宣布这药是真是假。")

    h(doc, "第5页　屏上的两个状态")
    para(doc, "念：左边这张，是翻开之后。绿灯亮，屏上写 VERIFY、AIM 2S，提醒老人对准包装，大约两秒后自动拍。右边是合上之后，屏上写 GUARD、TALK，提醒对着麦说话。我们故意只留大字和颜色，不让老人在胸前读一段说明书。")

    h(doc, "第6页　验真：先把字读下来")
    para(doc, "念：验真的第一步是把包装上的字抄下来。右图是子女看板：拍了一盒「药到病除」的药膏，读到的主要是品名。左下是串口，心跳已经回来，结论是「未见明显风险」。这句话只表示眼前没有特别夸张的承诺，不是给这支药做认证，也请老师不要理解成我们鉴定它是真药。")

    h(doc, "第7页　验真：对照讲座传单")
    para(doc, "念：同一套规则，换一张讲座传单。这张「极品太岁」上，系统读到了教授头衔、国家中医专家、治病除根、三个月行走自如、中风十年全家遭殃。出现这种话术，就应该升高风险，建议先不要付款、先问子女。绿和红对照的是话术，不是给产品盖章。")

    h(doc, "第8页　守护：电话打成高风险")
    para(doc, "念：合上证件是守护。左边是看板里的一次判断：来电自称公安局，说卡要冻结，立刻转到安全账户，还不许告诉儿子。关键词命中之后是高风险，建议先安抚、挂断后用通讯录回拨、不要转到所谓安全账户，必要时打 96110。右边是板子整屏变红，写 GUARD、HIGH。老人低头就能看见：先停手。")

    h(doc, "第9页　听不清就给橙灯")
    para(doc, "念：麦克风在板子侧面小孔上。免提没有贴紧、说得太轻，就不会假装听懂。屏上是橙底 CHECK，网页会写没听清。这是有意留出来的诚实状态：不清楚就不给绿灯。可以再听一次，也可以让子女在看板上把字改完再判断。")

    h(doc, "第10页　判断写在同一处")
    para(doc, "念：整条链路就三块。板子负责拍、听、按，以及用色块显示等级；云端用同一份 Skill 做读图、听写和规则判断；子女看板同步原文和建议。对错不在五个人脑子里各写一套。下面这行串口是心跳：板子在线，最近一次验真结论已经回到电脑。")

    h(doc, "第11页　分工对照表")
    para(doc, "念：这一页请老师看一下现在的分工。左边是已经做成的，右边是还可以精进的。杨立颖负责拍和读包装；马凤玲负责听和听不清时的提示；我负责验真、守护、求助这条状态机；贾喻然负责云端规则和看板；彭思羽负责把板子和电脑联在同一网络上。细处不展开，表上写的就是我们下一步还要打磨的地方。")

    h(doc, "第12页　现在与下一步")
    para(doc, "念：向老师交代现状：产品三条闭环——验真、守护、求助——已经在这块 ESP32-S3-EYE 上用 Arduino 跑通。赛题要求基于 openvela。应用和协议已经按同一份 HTTP 写在仓库里，还没有编烧到真机。所以不是还没做产品，而是产品演示已通，赛题真机对接是下一步。截止日期是 9 月 20 日。")

    h(doc, "第13页　谢谢")
    para(doc, "念：我们希望老人在没有子女陪同时，胸前仍有一个温和的校验点：先停手，再问清。以上是见微随身证的汇报，我是袁艺轩，谢谢老师，欢迎提问。")

    h(doc, "提问时怎么接（不投影）")
    para(doc, "若问真假药：我们只看话术和公开规则，不做药检，也不说国家认证。")
    para(doc, "若问为什么不是 openvela：现场为了把产品先跑稳，用的是 Arduino；openvela 路径已对齐，尚未烧录。")
    para(doc, "若点到个人：只请那位同学按表上「已经做成 / 还可以精进」各说一句。")
    para(doc, "若问板子没喇叭：原版开发板没有扬声器，等级主要靠屏上色块；网页不再自动朗读。")

    out = ROOT / "docs" / "见微随身证_最终汇报口播稿.docx"
    doc.save(out)
    return out


if __name__ == "__main__":
    photos = collect_photos()
    ppt = make_ppt(photos)
    speech = make_speech()
    print(ppt)
    print(speech)
    print("pages", TOTAL)
