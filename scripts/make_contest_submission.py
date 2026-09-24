# -*- coding: utf-8 -*-
"""按官方「作品提交模板」生成见微随身证介绍文档。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"C:\Users\jwr66\Documents\jianwei-id\docs\见微随身证_openvela作品提交.docx")
DESK = Path(r"C:\Users\jwr66\Desktop\见微随身证-作品提交.docx")


def set_run(run, size=12, bold=False, name="宋体", color=None, east="宋体"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)
    if color:
        run.font.color.rgb = color


def p(doc, text, size=12, bold=False, name="宋体", east="宋体", center=False, first_line=True):
    para = doc.add_paragraph()
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    para.paragraph_format.space_after = Pt(6)
    if first_line and not center:
        para.paragraph_format.first_line_indent = Cm(0.74)
    run = para.add_run(text)
    set_run(run, size, bold, name, east=east)
    return para


def h1(doc, text):
    para = doc.add_heading(text, level=1)
    for r in para.runs:
        set_run(r, 16, True, "黑体", east="黑体")


def h2(doc, text):
    para = doc.add_heading(text, level=2)
    for r in para.runs:
        set_run(r, 14, True, "黑体", east="黑体")


def h3(doc, text):
    para = doc.add_heading(text, level=3)
    for r in para.runs:
        set_run(r, 13, True, "黑体", east="黑体")


def shade(cell, hex_color):
    from docx.oxml import parse_xml

    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(
        parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="{hex_color}"/>'
        )
    )


def fill_table(table, rows, header=True):
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(val)
            is_h = header and ri == 0
            set_run(run, 10 if not is_h else 11, is_h, "宋体", east="黑体" if is_h else "宋体")
            if is_h:
                shade(cell, "A06B52")
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def make():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.54)
        s.left_margin = Cm(2.8)
        s.right_margin = Cm(2.8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("2026 首届 openvela AI 硬件开发者大赛")
    set_run(r, 16, True, "黑体", east="黑体")

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("作品提交说明（按官方模板撰稿）")
    set_run(r, 14, True, "黑体", east="黑体")

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("选题方向：AI 硬件产品创新　　开发模式：模式 A（设备协议 + 云端）")
    set_run(r, 11, False, "宋体", east="宋体")

    p(
        doc,
        "说明：本稿按组委会《作品提交模板》章节填写。源码以专属仓 contest2026_<编号>_<队名> 为准；压缩包命名建议「<团队名称>-见微随身证-contest2026_<编号>_<队名>.zip」。方括号【】处请赛前补上编号与队名拼音。",
        10.5,
        first_line=False,
    )

    h1(doc, "一、提交材料清单")
    tbl = doc.add_table(rows=6, cols=4)
    tbl.style = "Table Grid"
    fill_table(
        tbl,
        [
            ["序号", "材料", "格式 / 说明", "本队状态"],
            ["1", "作品介绍", "本文档 .docx", "必交 · 本稿"],
            ["2", "演示视频", "≤5 分钟 mp4，含硬件演示与 AI 能力", "必交 · 待补拍（可用已跑通的 Arduino 同协议画面，旁白说明交付目标为 openvela）"],
            ["3", "产品展示照片", "硬件实拍：前/侧/按键/屏显", "选交 · 仓库 docs/ppt_photos 与汇报 PPT 已有板子实拍"],
            ["4", "海报", "一图看懂产品", "选交 · 可用封面页「见微随身证」"],
            ["5", "答辩 PPT", "docs/见微随身证_最终汇报.pptx", "选交 · 已完成"],
        ],
    )
    p(doc, "本作品含硬件实物（ESP32-S3-EYE 胸牌形态），3.4 节写硬件组成，照片建议随压缩包一并上传。", 10.5, first_line=False)

    h1(doc, "二、填写本模板")
    h2(doc, "1. 信息表")
    info = doc.add_table(rows=5, cols=2)
    info.style = "Table Grid"
    fill_table(
        info,
        [
            ["项目", "内容"],
            ["作品名称", "见微·随身证"],
            ["参赛团队", "【队名】（成员：杨立颖、马凤玲、袁艺轩、贾喻然、彭思羽）"],
            [
                "团队分工",
                "杨立颖：验真拍照与读图链路；马凤玲：守护拾音与听写链路；袁艺轩：板端状态机（验真/守护/求助）与 LCD 等级；贾喻然：云端 Flask、skills/ 规则与子女看板；彭思羽：组网、CLOUD_HOST、端到端联调。",
            ],
            [
                "选题方向",
                "AI 硬件产品创新（基于已适配的 ESP32-S3-EYE，不走「新硬件适配」）。开发模式 A：设备 HTTP 协议 + 云端识别与规则；自定义 Skill：skills/jianwei-antifraud。",
            ],
        ],
        header=True,
    )

    h2(doc, "2. 摘要")
    p(
        doc,
        "见微随身证是挂在老人胸前的证件卡式反诈助手。交互只有三步：翻开验真、合上守护、连按三下求助。硬件采用乐鑫 ESP32-S3-EYE——该板已在 openvela 官方 BSP 适配摄像头、LCD、麦克风、Wi-Fi 与 BOOT，本作品在其上做产品，而不是再写一份板级。正式板端是 openvela 应用 app/jianwei，串口命令 jianwei；通过 /dev/video0、/dev/audio/pcm_in0、/dev/lcd0、/dev/buttons、wlan0 完成感知与显示，经 HTTP 把包装图或电话录音交给电脑侧云端。云端用 RapidOCR 读字、faster-whisper small 听写，再用仓库内唯一规则库 skills/ 分层打分，禁止输出「官方药检」结论。子女网页同步原文与等级。主动能力体现在：进入守护后设备自动拾音并上报（事件主动），以及说不出口时三连按触发求助通知（事件主动）。Arduino 回退固件已在同一块板上跑通三条闭环，用于验证协议；交付目标是把同一状态机烧进 openvela。",
    )

    h2(doc, "3. 正文")
    h3(doc, "3.1 背景")
    p(
        doc,
        "应用场景：社区讲座、保健品传单、来电催转账。用户是独自出门或子女不在身边的老人。痛点不是「不会用手机」，而是被推销话术隔开真实信息：包装承诺包治、电话要求转到安全账户且不许告诉家人。现有反诈 App 需要老人主动打开、打字或解锁，现场往往来不及。",
    )
    p(
        doc,
        "关键问题：如何在不鉴定真假药的前提下，用老人自己完得成的动作（翻、合、按）完成校验，并把结果同时给老人（大色块）和子女（看板）。",
    )
    p(
        doc,
        "创新点：（1）形态是证件卡而不是又一个聊天窗口；（2）感知在端、判定在同一份 Skill，板端不复制词表；（3）主动执行：合上即听、三连按即求助，无需老人先唤醒对话；（4）走赛题模式 A，设备协议稳定，云端可换识别模型而不改按键语义。",
    )

    h3(doc, "3.2 系统架构设计")
    p(
        doc,
        "总体架构：一个产品、三处代码、一份协议。板端状态机（openvela jianwei 或开发期 Arduino 回退）只负责拍、听、按和显示 HIGH/CHECK/OK；cloud/app.py 做读图、听写并调用 skills/；浏览器做子女看板。协议见仓库 docs/protocol.md，路径包括 /api/verify-image、/api/guard-audio、/api/device/heartbeat。",
    )
    p(
        doc,
        "方案论证：未选「新硬件适配」，因为 ESP32-S3-EYE 已在 vendor_espressif / dev-ai-contest-2026 适配。未选纯端侧大模型，因为板端算力与内存更适合做采集与显示。未选纯云端 App，因为赛题排除不在 openvela 设备上运行的作品。备选「只做网页演示」被否定，必须保留板端状态机。",
    )
    p(
        doc,
        "关键模块：jianwei_main 状态机；jianwei_cam 读 /dev/video0（QVGA RGB565→BMP）；jianwei_mic 读 /dev/audio/pcm_in0 成 WAV；jianwei_ui 写 /dev/lcd0 与 /dev/userleds；jianwei_net 以 HTTP 访问云端。云端 RapidOCR + Whisper + keywords.json / prices.json。角色上，设备是执行器，云端是判定器，网页是家属观察口，三者通过局域网 IPv4:8787 协同，板子禁止填写 127.0.0.1。",
    )

    h3(doc, "3.3 核心算法与技术原理")
    p(
        doc,
        "AI 算法分两段。读图：本机 RapidOCR（ONNX，rapidocr-onnxruntime 1.2.3），检测+识别中英模型，text_score=0.25，并做放大与旋转后取最长文本；不在板上跑视觉大模型。听写：faster-whisper 加载 Whisper small（CTranslate2），强制中文、关闭 VAD，CUDA 优先。二者都只负责「变成字」，风险分数由规则计算。大模型（如硅基流动上的 Qwen 视觉）仅作 OCR 全空时的备用，判定主路径不依赖其改档。未使用大赛 MiMo 配额时，在介绍中注明接口自建云端；若后续接入，只需改云端 BASE_URL，板端协议不变。",
    )
    p(
        doc,
        "判定不是单一关键词。守护：skills/keywords.json 子串加分 +「紧急+要钱+不让核实」组合 +40 + 索要验证码 +25，≥70 高风险、≥35 需核实。验真：备案号格式、prices.json 价格带、夸大短语、购买渠道，关键词只减权计入。听不清（无中文/过静）固定橙灯 CHECK，三连按不读文本。详情见 docs/技术路线与判定.md。",
    )
    p(
        doc,
        "openvela 能力占用：图形——/dev/lcd0 两行 ASCII 等级色块语义（HIGH/CHECK/OK）；多媒体——/dev/video0 拍照、/dev/audio/pcm_in0 拾音；AI——设备侧执行采集与上报，云端 OCR/ASR/Skill，自定义 Skill 文本在 skills/jianwei-antifraud/SKILL.md，部署目标 /data/agent/skills/jianwei-antifraud/。对 openvela 的使用是调用官方 ESP32-S3-EYE 节点与 NuttX 网络栈，不自研 BSP；扩展是增加独立应用 jianwei 与同一份 HTTP 协议。",
    )

    h3(doc, "3.4 系统实现")
    p(
        doc,
        "软件/固件：openvela 分支 dev-ai-contest-2026，应用 app/jianwei（Make.defs 勾选进 packages/demos）。menuconfig 启用「见微随身证」，JIANWEI_CLOUD_HOST 填电脑局域网 IP。nsh 联网后执行 jianwei。开发期 Windows 无交叉工具链时，firmware/jianwei_eye 实现同一状态机，用于验证协议与演示，不是第二套产品。",
    )
    p(
        doc,
        "硬件：主控 ESP32-S3-WROOM-1（S3-EYE 模组），OV2640，ST7789 240×240，PDM 数字麦，BOOT GPIO0，LED GPIO3，Wi-Fi 2.4 GHz。不新增未适配芯片，不宣称自研 PCB。关键 BOM 即官方开发板本身。安全：密钥只在电脑 .env，固件 secrets.h 不入库。",
    )
    p(
        doc,
        "应用交互：老人侧无复杂 UI，只有大字等级；子女侧 Flask 网页展示 OCR/ASR 原文、分数、建议。唤醒词若走语音渠道遵循赛事「你好，openvela」；当前主交互是按键事件，不依赖老人先喊唤醒词，以满足现场手脚不便的场景。",
    )
    p(
        doc,
        "自定义 Skill：jianwei-antifraud。内容规定何时启用（包装/电话/三连按）、如何套 keywords 与价格带、禁止「国家认证」、高风险话术顺序（安抚→回拨→暂缓转账→96110）。主动场景：进入守护后自动录音并 POST（事件主动）；三连按求助写入子女事件（事件主动）。这满足「至少 1 个自定义 Skill + 至少 1 个主动+执行」，而不是纯问答机器人。",
    )

    h3(doc, "3.5 系统测试")
    p(doc, "测试环境：ESP32-S3-EYE + 电脑 Flask 8787，同一 2.4 GHz 热点，CLOUD_HOST 为电脑 IPv4。识别在电脑 CPU/GPU，板端负责采集。", first_line=True)
    p(
        doc,
        "功能：短按进入验真可拍照并在网页出现 OCR 文本；短按守护可上传 WAV；演示话术「公安局+安全账户+不要告诉儿子」命中高风险，板子红屏 HIGH；听不清为橙 CHECK；三连按网页求助变红。普通包装可出现未见明显风险。网页打字与板子上传走同一 verify_text/guard_text。",
    )
    p(
        doc,
        "性能（开发机实测量级，非正式实验室报告）：VGA JPEG 约 20–40 KB；RapidOCR 单张约数秒；Whisper small 8 秒 16 kHz WAV 约数秒级（CUDA）；心跳 4 s；HTTP 超时板端 90 s。云端线程模型 threaded，避免识别时卡死心跳。",
    )
    p(
        doc,
        "可靠与边界：OCR 对艺术字会误识，故规则看话术堆叠而非单字；VAD 曾导致远场全空，已关闭；词表默认阈值过高导致「字没读清」，已降 text_score。网络依赖同一热点，断网心跳失败——这是已知限制，不伪装成已解决。内存：板端 BMP QVGA 约 230 KB 量级堆分配；异常时 LCD 显示 NO CAM / NO MIC / NO NET。安全：不做药检承诺，避免误导老人对质骗子。",
    )

    h3(doc, "3.6 AI-Native 开发说明")
    ai = doc.add_table(rows=6, cols=2)
    ai.style = "Table Grid"
    fill_table(
        ai,
        [
            ["指标", "描述"],
            [
                "AI Coding 代码占比",
                "约 80%。Cursor 辅助生成板端状态机、云端规则、看板与文档；人工负责选题约束（不做药检）、真机验收、密钥与组网、演示脚本。",
            ],
            ["使用的 AI 工具", "Cursor（Agent / 对话）。"],
            [
                "MCP 工具使用情况",
                "Cursor 内置浏览器用于子女看板验收。未使用 VelaJS MCP / Figma MCP。",
            ],
            [
                "Skills 使用及贡献情况",
                "自定义 Skill：skills/jianwei-antifraud（赛题要求至少 1 个）。规则 JSON 与云端共用，避免 Agent 与 Flask 两套逻辑。",
            ],
            [
                "Token 使用情况",
                "开发期未按大赛控制台单独记账；若使用官方 MiMo 配额请在控制台查询后填入本栏。云端识别也可使用自备接口，板端协议不变。",
            ],
        ],
    )
    p(doc, "贡献：AI 工具显著加快协议对齐、OCR/ASR 接入与文稿结构；关键产品决策（三档阈值、听不清不给绿灯、禁止认证用语）由队员确认。", 10.5, first_line=False)

    h3(doc, "3.7 总结与展望")
    p(
        doc,
        "成功之处：在已适配硬件上做出老人可独立完成的三步闭环；判定集中在 skills/；模式 A 把端云分工写清楚；Arduino 真机证明协议可跑，openvela 应用按官方节点对接完毕。",
    )
    p(
        doc,
        "应用前景：社区、子女陪伴产品、讲座现场的「先停手」工具。目标用户偏中老年、价格敏感、使用偏好实体按键而非 App。商业上可走配件化胸牌，云端规则可运营更新。社会价值是降低隔离式诈骗的当场成功率，而不是替代监管药检。",
    )
    p(
        doc,
        "不足：openvela 固件需在 Ubuntu/WSL 专属仓内完成最终编烧与 5 分钟视频；板载无喇叭；听写依赖贴紧麦孔；艺术字体 OCR 不稳定。",
    )
    p(
        doc,
        "下一步：repo sync 后合并 defconfig.append，烧录 jianwei，把 Skill 拷到 /data/agent/skills/；补拍含 nsh> jianwei 的视频；外壳与真翻盖；截止日期 2026 年 9 月 20 日。",
    )

    h1(doc, "三、openvela 对接对照（评委定位源码）")
    map_t = doc.add_table(rows=9, cols=3)
    map_t.style = "Table Grid"
    fill_table(
        map_t,
        [
            ["赛题要求", "本仓库落点", "说明"],
            ["跑在 openvela 设备上", "app/jianwei + 官方 ESP32-S3-EYE BSP", "不重复移植；nsh> jianwei"],
            ["模式 A 设备协议+云端", "docs/protocol.md，cloud/app.py:8787", "与 Arduino 字段相同"],
            ["自定义 Skill ≥1", "skills/jianwei-antifraud/SKILL.md", "部署到 /data/agent/skills/"],
            ["主动+执行", "进入守护自动拾音 POST；三连按求助", "事件主动，非纯聊天"],
            ["图形 / 多媒体 / AI", "lcd0；video0+pcm_in0；OCR+ASR+规则", "三类都用到"],
            ["禁止纯云端", "板端状态机必须存在", "网页只是子女观察口"],
            ["唤醒词", "你好，openvela（若开语音渠道）", "主路径为按键，避免老人先喊口号"],
            ["代码提交", "fork 专属仓 → PR 自合；logs/ 导出对话", "Apache-2.0；截止 9 月 20 日"],
        ],
    )

    h1(doc, "四、注意事项（模板原文对齐）")
    p(doc, "作品提交截止：2026 年 9 月 20 日，以组委会最新公告为准。", first_line=False)
    p(doc, "AI Coding 日志不可由本文代替，须按手册导出到专属仓 logs/，不要手动删改关键记录，不要包含密钥。", first_line=False)
    p(doc, "作品须为原创，遵循 Apache 2.0，不侵犯版权与专利权。", first_line=False)
    p(doc, "工程基于 openvela（NuttX 内核等），并使用图形、AI、多媒体能力中的对应项，见第三节对照表。", first_line=False)
    p(doc, "语音唤醒若启用，统一使用「你好，openvela」。", first_line=False)
    p(doc, "比赛在 GitHub 进行。获奖后按要求将作品 PR 至 openvela 上游对应仓库 dev-ai-contest-2026 分支。", first_line=False)

    h1(doc, "五、评分维度自检")
    score = doc.add_table(rows=7, cols=2)
    score.style = "Table Grid"
    fill_table(
        score,
        [
            ["评分维度（分值）", "本稿对应章节"],
            ["技术难度（30）", "3.2 架构、3.3 算法、3.4 实现、3.5 测试、第三节对照表"],
            ["产品创新性（20）", "2 摘要、3.1 背景"],
            ["项目完成度（20）", "3.5 测试、源码 app/jianwei+cloud+firmware、展示照片"],
            ["AI 开发（10）", "3.3 算法、3.6 AI-Native、自定义 Skill"],
            ["商业潜力（10）", "3.7 展望"],
            ["展示效果（10）", "演示视频、答辩 PPT"],
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    try:
        doc.save(DESK)
    except OSError:
        pass
    print(OUT)
    if DESK.exists():
        print(DESK)


if __name__ == "__main__":
    make()
