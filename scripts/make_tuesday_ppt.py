# -*- coding: utf-8 -*-
"""Generate a warm, keyword-light Tuesday briefing deck."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from lxml import etree
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# 16:9
W, H = Inches(13.333), Inches(7.5)

CREAM = RGBColor(0xFB, 0xF4, 0xEC)
PAPER = RGBColor(0xFF, 0xFB, 0xF6)
INK = RGBColor(0x4A, 0x3A, 0x30)
SOFT = RGBColor(0x8A, 0x73, 0x64)
BROWN = RGBColor(0xA0, 0x6B, 0x52)
CORAL = RGBColor(0xC4, 0x6B, 0x5A)
SAGE = RGBColor(0x6A, 0x8F, 0x72)
GOLD = RGBColor(0xC9, 0xA0, 0x68)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run(run, size, bold=False, color=INK, font="微软雅黑"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    r = run._r
    rPr = r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font)
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", font)


def fill_slide(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, l, t, w, h, color, radius=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    if radius is not None:
        shape.adjustments[0] = radius
    else:
        shape.adjustments[0] = 0.08
    return shape


def add_text(slide, l, t, w, h, text, size, bold=False, color=INK, align=PP_ALIGN.LEFT, font="微软雅黑"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color, font)
    return box


def add_lines(slide, l, t, w, h, lines, size, bold=False, color=INK, align=PP_ALIGN.LEFT, gap=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(gap)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return box


def footer(slide, page, total=11):
    add_text(
        slide,
        Inches(0.6),
        Inches(7.1),
        Inches(8),
        Inches(0.3),
        "见微 · 随身证",
        11,
        color=SOFT,
    )
    add_text(
        slide,
        Inches(11.2),
        Inches(7.1),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        11,
        color=SOFT,
        align=PP_ALIGN.RIGHT,
    )


def accent_bar(slide):
    add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.12), BROWN, radius=0)


def make():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]

    # 1 cover
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    add_rect(s, Inches(4.15), Inches(1.15), Inches(5.05), Inches(3.15), PAPER, 0.12)
    add_rect(s, Inches(4.15), Inches(1.15), Inches(5.05), Inches(0.18), BROWN, 0)
    add_text(s, Inches(4.2), Inches(1.55), Inches(5), Inches(0.4), "胸前的证件卡", 16, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(4.2), Inches(2.0), Inches(5), Inches(0.9), "见微", 54, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(4.2), Inches(2.95), Inches(5), Inches(0.5), "随  身  证", 28, True, BROWN, PP_ALIGN.CENTER)
    add_text(s, Inches(2), Inches(4.7), Inches(9.3), Inches(0.6), "翻开验真    ·    合上守护", 26, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(2), Inches(5.5), Inches(9.3), Inches(0.4), "先看清话术，先停手", 16, color=SOFT, align=PP_ALIGN.CENTER)

    # 2 why
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 2)
    add_text(s, Inches(0.8), Inches(1.4), Inches(12), Inches(0.5), "老人最怕的", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(2.1), Inches(12), Inches(1.2), "不是不会用手机", 36, True, INK)
    add_text(s, Inches(0.8), Inches(3.4), Inches(12), Inches(1.0), "是被隔绝真实信息", 36, True, BROWN)
    add_text(s, Inches(0.8), Inches(5.0), Inches(12), Inches(0.5), "讲座  ·  包装  ·  电话", 18, color=SOFT)

    # 3 what
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 3)
    add_text(s, Inches(0.8), Inches(1.3), Inches(12), Inches(0.5), "我们做的", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(1.9), Inches(12), Inches(0.8), "挂在胸前的校验点", 32, True, INK)
    cards = [
        (0.8, "看得见", "包装  ·  传单"),
        (4.9, "听得到", "电话  ·  讲座"),
        (9.0, "说不出", "连按三下"),
    ]
    for x, title, sub in cards:
        add_rect(s, Inches(x), Inches(3.3), Inches(3.5), Inches(2.4), PAPER, 0.1)
        add_text(s, Inches(x), Inches(3.7), Inches(3.5), Inches(0.7), title, 26, True, BROWN, PP_ALIGN.CENTER)
        add_text(s, Inches(x), Inches(4.55), Inches(3.5), Inches(0.6), sub, 16, color=SOFT, align=PP_ALIGN.CENTER)

    # 4 how
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 4)
    add_text(s, Inches(0.8), Inches(1.2), Inches(12), Inches(0.5), "怎么用", 16, color=SOFT)
    rows = [
        (1.9, BROWN, "翻开", "验真"),
        (3.55, SAGE, "合上", "守护"),
        (5.2, CORAL, "三连按", "求助"),
    ]
    for y, c, a, b in rows:
        add_rect(s, Inches(2.2), Inches(y), Inches(8.9), Inches(1.35), PAPER, 0.1)
        add_rect(s, Inches(2.2), Inches(y), Inches(0.18), Inches(1.35), c, 0)
        add_text(s, Inches(2.8), Inches(y + 0.28), Inches(3.5), Inches(0.8), a, 32, True, INK)
        add_text(s, Inches(7.0), Inches(y + 0.35), Inches(3.5), Inches(0.7), b, 28, True, c)

    # 5 demo verify ok
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 5)
    add_text(s, Inches(0.8), Inches(1.2), Inches(12), Inches(0.4), "现场  ·  验真", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(1.8), Inches(12), Inches(0.9), "拍一拍普通包装", 36, True, INK)
    add_rect(s, Inches(0.8), Inches(3.3), Inches(5.5), Inches(2.5), PAPER, 0.1)
    add_text(s, Inches(0.8), Inches(3.7), Inches(5.5), Inches(0.6), "应看到", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(4.3), Inches(5.5), Inches(0.9), "绿  ·  OK", 36, True, SAGE, PP_ALIGN.CENTER)
    add_rect(s, Inches(6.7), Inches(3.3), Inches(5.8), Inches(2.5), PAPER, 0.1)
    add_text(s, Inches(6.7), Inches(3.7), Inches(5.8), Inches(0.6), "请", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(6.7), Inches(4.3), Inches(5.8), Inches(0.9), "杨立颖 同学", 28, True, BROWN, PP_ALIGN.CENTER)

    # 6 demo high risk
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 6)
    add_text(s, Inches(0.8), Inches(1.2), Inches(12), Inches(0.4), "现场  ·  验真", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(1.8), Inches(12), Inches(0.9), "讲座里的推销话术", 36, True, INK)
    add_rect(s, Inches(0.8), Inches(3.3), Inches(5.5), Inches(2.5), PAPER, 0.1)
    add_text(s, Inches(0.8), Inches(3.7), Inches(5.5), Inches(0.6), "应看到", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(4.3), Inches(5.5), Inches(0.9), "红  ·  先停手", 32, True, CORAL, PP_ALIGN.CENTER)
    add_rect(s, Inches(6.7), Inches(3.3), Inches(5.8), Inches(2.5), PAPER, 0.1)
    add_text(s, Inches(6.7), Inches(3.7), Inches(5.8), Inches(0.6), "请", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(6.7), Inches(4.3), Inches(5.8), Inches(0.9), "贾喻然 同学", 28, True, BROWN, PP_ALIGN.CENTER)

    # 7 demo guard
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 7)
    add_text(s, Inches(0.8), Inches(1.2), Inches(12), Inches(0.4), "现场  ·  守护", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(1.8), Inches(12), Inches(0.9), "听一听电话话术", 36, True, INK)
    add_text(s, Inches(0.8), Inches(2.7), Inches(12), Inches(0.4), "公安局    安全账户    不要告诉家人", 16, color=SOFT)
    add_rect(s, Inches(0.8), Inches(3.5), Inches(5.5), Inches(2.3), PAPER, 0.1)
    add_text(s, Inches(0.8), Inches(3.85), Inches(5.5), Inches(0.5), "应看到", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(4.4), Inches(5.5), Inches(0.8), "红  ·  HIGH", 32, True, CORAL, PP_ALIGN.CENTER)
    add_rect(s, Inches(6.7), Inches(3.5), Inches(5.8), Inches(2.3), PAPER, 0.1)
    add_text(s, Inches(6.7), Inches(3.85), Inches(5.8), Inches(0.5), "请", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(6.7), Inches(4.4), Inches(5.8), Inches(0.8), "马凤玲 同学", 28, True, BROWN, PP_ALIGN.CENTER)

    # 8 help
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 8)
    add_text(s, Inches(0.8), Inches(1.2), Inches(12), Inches(0.4), "现场  ·  求助", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(1.8), Inches(12), Inches(0.9), "连按三下", 36, True, INK)
    add_rect(s, Inches(0.8), Inches(3.3), Inches(5.5), Inches(2.5), PAPER, 0.1)
    add_text(s, Inches(0.8), Inches(3.7), Inches(5.5), Inches(0.6), "应看到", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(4.3), Inches(5.5), Inches(0.9), "HELP  ·  红点", 32, True, CORAL, PP_ALIGN.CENTER)
    add_rect(s, Inches(6.7), Inches(3.3), Inches(5.8), Inches(2.5), PAPER, 0.1)
    add_text(s, Inches(6.7), Inches(3.7), Inches(5.8), Inches(0.6), "请", 14, color=SOFT, align=PP_ALIGN.CENTER)
    add_text(s, Inches(6.7), Inches(4.3), Inches(5.8), Inches(0.9), "袁艺轩 同学", 28, True, BROWN, PP_ALIGN.CENTER)

    # 9 tech
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 9)
    add_text(s, Inches(0.8), Inches(1.3), Inches(12), Inches(0.5), "怎么判断", 16, color=SOFT)
    add_text(s, Inches(0.8), Inches(1.95), Inches(12), Inches(0.7), "同一套规则", 32, True, INK)
    trio = [
        (0.8, "板子", "拍  ·  听  ·  按"),
        (4.9, "云端", "Skill 判断"),
        (9.0, "子女", "看板  ·  朗读"),
    ]
    for x, a, b in trio:
        add_rect(s, Inches(x), Inches(3.3), Inches(3.5), Inches(2.4), PAPER, 0.1)
        add_text(s, Inches(x), Inches(3.75), Inches(3.5), Inches(0.7), a, 26, True, BROWN, PP_ALIGN.CENTER)
        add_text(s, Inches(x), Inches(4.6), Inches(3.5), Inches(0.6), b, 16, color=SOFT, align=PP_ALIGN.CENTER)

    # 10 now
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    footer(s, 10)
    add_text(s, Inches(0.8), Inches(1.4), Inches(12), Inches(0.5), "现在到哪一步", 16, color=SOFT)
    add_rect(s, Inches(0.8), Inches(2.3), Inches(5.7), Inches(3.5), PAPER, 0.1)
    add_text(s, Inches(0.8), Inches(2.85), Inches(5.7), Inches(0.45), "已经跑通", 16, color=SAGE, align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(3.45), Inches(5.7), Inches(0.7), "Arduino", 28, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(0.8), Inches(4.2), Inches(5.7), Inches(0.7), "三条闭环", 28, True, INK, PP_ALIGN.CENTER)
    add_rect(s, Inches(6.9), Inches(2.3), Inches(5.6), Inches(3.5), PAPER, 0.1)
    add_text(s, Inches(6.9), Inches(2.85), Inches(5.6), Inches(0.45), "下一步", 16, color=GOLD, align=PP_ALIGN.CENTER)
    add_text(s, Inches(6.9), Inches(3.45), Inches(5.6), Inches(0.7), "openvela", 28, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(6.9), Inches(4.2), Inches(5.6), Inches(0.7), "真机对接", 28, True, INK, PP_ALIGN.CENTER)

    # 11 thanks
    s = prs.slides.add_slide(blank)
    fill_slide(s, CREAM)
    accent_bar(s)
    add_text(s, Inches(1), Inches(2.3), Inches(11.3), Inches(1.0), "先停手，再问清", 36, True, INK, PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(3.5), Inches(11.3), Inches(0.6), "谢谢老师", 22, color=BROWN, align=PP_ALIGN.CENTER)
    add_text(s, Inches(1), Inches(5.2), Inches(11.3), Inches(0.4), "杨立颖  ·  马凤玲  ·  袁艺轩  ·  贾喻然  ·  彭思羽", 14, color=SOFT, align=PP_ALIGN.CENTER)

    out = Path(r"C:\Users\jwr66\Documents\jianwei-id\docs\见微随身证_周二汇报.pptx")
    prs.save(out)
    print(out)


if __name__ == "__main__":
    make()
